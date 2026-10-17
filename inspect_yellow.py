import docx
import xml.etree.ElementTree as ET

doc = docx.Document('Claude_MediScan_revised.docx')
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
w_fill = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill'

yellow_items = []

# Search paragraphs
for p_idx, p in enumerate(doc.paragraphs):
    # Check pPr shading
    pPr = p._p.find('w:pPr', ns)
    if pPr is not None:
        shd = pPr.find('w:shd', ns)
        if shd is not None and shd.get(w_fill) == 'FFFF00':
            yellow_items.append(('paragraph_shd', p_idx, None, p.text, p.text))
    
    # Check run shading
    for r_idx, r in enumerate(p.runs):
        rPr = r._r.find('w:rPr', ns)
        if rPr is not None:
            shd = rPr.find('w:shd', ns)
            if shd is not None and shd.get(w_fill) == 'FFFF00':
                yellow_items.append(('run_shd', p_idx, r_idx, r.text, p.text))

# Search tables
for t_idx, t in enumerate(doc.tables):
    for r_idx, row in enumerate(t.rows):
        for c_idx, cell in enumerate(row.cells):
            tcPr = cell._tc.find('w:tcPr', ns)
            if tcPr is not None:
                shd = tcPr.find('w:shd', ns)
                if shd is not None and shd.get(w_fill) == 'FFFF00':
                    yellow_items.append(('cell_shd', t_idx, (r_idx, c_idx), cell.text, cell.text))
            for cp_idx, p in enumerate(cell.paragraphs):
                for run_idx, r in enumerate(p.runs):
                    rPr = r._r.find('w:rPr', ns)
                    if rPr is not None:
                        shd = rPr.find('w:shd', ns)
                        if shd is not None and shd.get(w_fill) == 'FFFF00':
                            yellow_items.append(('table_run_shd', t_idx, (r_idx, c_idx, cp_idx, run_idx), r.text, cell.text))

print(f"Total yellow items found: {len(yellow_items)}")
for i, item in enumerate(yellow_items):
    print(f"[{i+1}] Type: {item[0]}, Index: {item[1]}, SubIndex: {item[2]}")
    print(f"    Target Text: {repr(item[3])}")
    print(f"    Context: {repr(item[4][:120])}")
    print("-" * 60)
